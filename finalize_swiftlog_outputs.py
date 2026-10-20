"""Regenerate SwiftLog Week 5 documentation and notebook from saved model outputs."""
from pathlib import Path
import json
import sys

import numpy as np
import pandas as pd
import joblib

ROOT = Path(__file__).resolve().parent
MODELS = ROOT / "Models"
DATA = ROOT / "Data"
PROCESSED = DATA / "Processed_Datasets"
REPORTS = ROOT / "Reports"
PRESENTATION = ROOT / "Presentation"
NOTEBOOK = ROOT / "Notebook"

for directory in (REPORTS, PRESENTATION, NOTEBOOK):
    directory.mkdir(parents=True, exist_ok=True)

comparison_path = MODELS / "Model_Comparison.xlsx"
bundle_path = MODELS / "Best_Model.pkl"

if not comparison_path.exists():
    raise FileNotFoundError(f"Required model comparison file not found: {comparison_path}")

if not bundle_path.exists():
    raise FileNotFoundError(f"Required saved model bundle not found: {bundle_path}")

comparison = pd.read_excel(comparison_path)
quality_path = PROCESSED / "data_quality_summary.csv"
quality = pd.read_csv(quality_path) if quality_path.exists() else pd.DataFrame()
bundles = joblib.load(bundle_path)

required_columns = {"Problem", "Algorithm"}
missing_columns = required_columns.difference(comparison.columns)
if missing_columns:
    raise ValueError(
        "Model_Comparison.xlsx is missing required columns: "
        + ", ".join(sorted(missing_columns))
    )

# Standardize selected/tuned flags where these columns exist.
for column in ("Selected", "Tuned"):
    if column in comparison.columns:
        comparison[column] = (
            comparison[column]
            .astype(str)
            .str.strip()
            .str.lower()
            .map({"true": True, "false": False, "1": True, "0": False})
            .fillna(False)
            .astype(bool)
        )

# Keep the original task order.
problem_order = [
    "Delivery Delay Classification",
    "Product Return Classification",
    "Product Demand Forecasting",
    "Customer Value Prediction (Next Month)",
]

comparison["_problem_order"] = (
    comparison["Problem"].astype(str).map(
        {name: index for index, name in enumerate(problem_order)}
    ).fillna(len(problem_order))
)

sort_columns = ["_problem_order"]
ascending = [True]

if "Selected" in comparison.columns:
    sort_columns.append("Selected")
    ascending.append(False)

sort_columns.append("Algorithm")
ascending.append(True)

comparison = comparison.sort_values(sort_columns, ascending=ascending)
comparison = comparison.drop(columns=["_problem_order"])
comparison.to_excel(comparison_path, index=False)

# ---------------------------------------------------------------------
# Metric helpers: support both classification and regression safely.
# ---------------------------------------------------------------------
def metric_value(row, metric):
    """Return a numeric metric, accepting common alternative column names."""
    possible_names = (
        metric,
        f"Test_{metric}",
        f"Heldout_{metric}",
        f"Held_Out_{metric}",
    )

    for name in possible_names:
        if name in row.index:
            value = row.get(name)
            if pd.notna(value):
                try:
                    return float(value)
                except (TypeError, ValueError):
                    pass

    return None


def format_metric(value, digits=3):
    return "not available" if value is None else f"{value:.{digits}f}"


def metric_line(task):
    """Create a task-specific metric summary without assuming metric columns."""
    rows = comparison[comparison["Problem"].astype(str) == task]

    if rows.empty:
        return f"{task}: metrics not available."

    selected_rows = (
        rows[rows["Selected"]]
        if "Selected" in rows.columns
        else rows.iloc[0:0]
    )
    row = selected_rows.iloc[0] if not selected_rows.empty else rows.iloc[0]

    algorithm = str(row.get("Algorithm", "Algorithm not recorded"))

    classification_tasks = {
        "Delivery Delay Classification",
        "Product Return Classification",
    }

    if task in classification_tasks:
        f1 = metric_value(row, "F1")
        auc = metric_value(row, "ROC_AUC")
        accuracy = metric_value(row, "Accuracy")
        precision = metric_value(row, "Precision")
        recall = metric_value(row, "Recall")

        parts = [f"Selected: {algorithm}"]

        if accuracy is not None:
            parts.append(f"Accuracy {accuracy:.3f}")
        if precision is not None:
            parts.append(f"Precision {precision:.3f}")
        if recall is not None:
            parts.append(f"Recall {recall:.3f}")
        if f1 is not None:
            parts.append(f"F1 {f1:.3f}")
        if auc is not None:
            parts.append(f"ROC-AUC {auc:.3f}")

        if len(parts) == 1:
            return f"{task}: classification metrics not available."

        return " | ".join(parts)

    mae = metric_value(row, "MAE")
    rmse = metric_value(row, "RMSE")
    r2 = metric_value(row, "R2")

    parts = [f"Selected: {algorithm}"]

    if mae is not None:
        parts.append(f"MAE {mae:.2f}")
    if rmse is not None:
        parts.append(f"RMSE {rmse:.2f}")
    if r2 is not None:
        parts.append(f"R-squared {r2:.3f}")

    if len(parts) == 1:
        return f"{task}: regression metrics not available."

    return " | ".join(parts)


# ---------------------------------------------------------------------
# Streamlit app text and warnings.
# ---------------------------------------------------------------------
app_path = ROOT / "App" / "app.py"

if app_path.exists():
    app = app_path.read_text(encoding="utf-8")

    app = app.replace('page_icon="Ã°Å¸Å¡Åš"', 'page_icon="📦"')
    app = app.replace(
        "Predicted daily product demand (units)",
        "Predicted weekly product demand (units)",
    )

    warning_block = '''if b["kind"] == "classification" and b["metrics"].get("ROC_AUC", 1) < 0.60:
    st.warning("The ROC-AUC indicates weak separation between classes. Treat this output as an experimental signal, not a reliable operational decision.")
if b["kind"] == "regression" and b["metrics"].get("R2", 0) < 0:
    st.warning("The held-out R-squared is below zero, so this model performed worse than a simple test-set mean baseline. Use the estimate for demonstration only.")'''

    metrics_line = 'st.write("Held-out test metrics:", b["metrics"])'

    if warning_block not in app and metrics_line in app:
        app = app.replace(metrics_line, metrics_line + "\n" + warning_block)

    app_path.write_text(app, encoding="utf-8")

# ---------------------------------------------------------------------
# Notebook.
# ---------------------------------------------------------------------
try:
    import nbformat as nbf

    nb = nbf.v4.new_notebook()
    cells = []

    cells.append(nbf.v4.new_markdown_cell("""# SwiftLog Week 5: Machine Learning Engineering

**Mission:** Intelligent Logistics Prediction & Supply Chain Optimization.

This notebook reviews source data, engineered features, model comparisons, evaluation artefacts, and limitations. The full model-training pipeline is `../run_swiftlog_project.py`. Run that script from the project root to regenerate the saved models and outputs before rerunning this notebook.

The supplied data is used as provided for this educational prototype. Results are reported as measured; weak performance is not hidden or described as production-ready."""))

    cells.append(nbf.v4.new_markdown_cell("## 1. Imports and paths"))

    cells.append(nbf.v4.new_code_cell("""from pathlib import Path
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

ROOT = Path('..').resolve()
DATA = ROOT / 'Data'
PROCESSED = DATA / 'Processed_Datasets'
MODELS = ROOT / 'Models'
"""))

    cells.append(nbf.v4.new_markdown_cell(
        "## 2. Source data and quality checks\n"
        "The source tables cover customers, orders, order items, products, "
        "shipments, delivery performance, returns, inventory, warehouses, "
        "suppliers, and employees. The supplied file-level data dictionary "
        "is brief; calculated model features are documented in "
        "`Feature_Data_Dictionary.csv`."
    ))

    cells.append(nbf.v4.new_code_cell("""quality_path = PROCESSED / 'data_quality_summary.csv'
if quality_path.exists():
    quality = pd.read_csv(quality_path)
    display(quality)
else:
    print('Data-quality summary is not available.')
"""))

    cells.append(nbf.v4.new_markdown_cell("## 3. Inspect the original target distributions"))

    cells.append(nbf.v4.new_code_cell("""original = DATA / 'Original_Datasets'
shipments = pd.read_csv(original / 'SwiftLog_Shipments.csv')
returns = pd.read_csv(original / 'SwiftLog_Returns.csv')
orders = pd.read_csv(original / 'SwiftLog_Orders.csv')

print('Shipments:', len(shipments))
if 'DelayDays' in shipments.columns:
    delays = pd.to_numeric(shipments['DelayDays'], errors='coerce')
    print('Delayed shipments (DelayDays > 0):', int((delays > 0).sum()))
if 'OrderID' in returns.columns:
    print('Orders associated with recorded returns:', returns['OrderID'].nunique())
if 'OrderID' in orders.columns:
    print('Orders:', orders['OrderID'].nunique())
"""))

    cells.append(nbf.v4.new_markdown_cell("""## 4. Engineered feature samples

The delivery-delay target is `DelayDays > 0`. The return target is whether an order appears in the Returns table. Demand features use weekly product quantities and lagged values from prior weeks. Customer value is framed as next-month revenue, using the preceding three months of activity. For the demand and customer-value tasks, evaluation is chronological rather than a random row split."""))

    feature_files = [
        ("Delivery delay features", "delivery_delay_features.csv"),
        ("Product return features", "product_return_features.csv"),
        ("Weekly demand forecasting features", "demand_forecasting_features.csv"),
        ("Customer next-month value features", "customer_value_features.csv"),
    ]

    for title, filename in feature_files:
        cells.append(nbf.v4.new_markdown_cell("### " + title))
        cells.append(nbf.v4.new_code_cell(
            f"""feature_path = PROCESSED / '{filename}'
if feature_path.exists():
    features = pd.read_csv(feature_path)
    print('Shape:', features.shape)
    display(features.head())
else:
    print('Feature file not found:', feature_path)
"""
        ))

    cells.append(nbf.v4.new_markdown_cell("""## 5. Compare the candidate algorithms

The table contains candidate algorithms per task. `Selected` marks the model stored in `Best_Model.pkl`; `Tuned` indicates that a tuning search was accepted, where recorded. The selected model metrics are stored in the comparison table."""))

    cells.append(nbf.v4.new_code_cell("""comparison = pd.read_excel(MODELS / 'Model_Comparison.xlsx')
display(comparison)
"""))

    cells.append(nbf.v4.new_markdown_cell("## 6. Review selected models and held-out metrics"))

    cells.append(nbf.v4.new_code_cell("""if 'Selected' in comparison.columns:
    selected = comparison[comparison['Selected'] == True]
    display(selected)
else:
    print('The comparison table does not contain a Selected column.')
"""))

    cells.append(nbf.v4.new_markdown_cell("## 7. Evaluation visuals"))

    cells.append(nbf.v4.new_code_cell("""from IPython.display import Image, display

for name in [
    'delivery_delay_classification_confusion_matrix.png',
    'product_return_classification_confusion_matrix.png',
]:
    path = MODELS / name
    if path.exists():
        print(name)
        display(Image(filename=str(path)))

importance = MODELS / 'Feature_Importance.png'
if importance.exists():
    display(Image(filename=str(importance)))
"""))

    cells.append(nbf.v4.new_markdown_cell("""## 8. Interpretation and limitations

- Delivery-delay and return-classification ROC-AUC values near 0.5 indicate weak class discrimination, even when F1 or accuracy appears high.
- Negative held-out R-squared for demand or next-month customer value means the selected model performed worse than a simple mean-prediction baseline on the test period.
- The dataset spans a limited period and has limited operational context. Forecast quality may improve with richer timestamps, planned delivery dates, inventory snapshots, and more history.
- Model outputs should support human review only. They are not validated for live logistics decisions.

## 9. Reproduce the outputs

From the project root, run `python run_swiftlog_project.py`. The script preprocesses the source data, builds the features, compares candidate algorithms, evaluates candidates, saves model bundles and visualizations, and regenerates reports and the presentation. Then rerun this notebook to inspect the refreshed outputs."""))

    nb.cells = cells
    nb["metadata"] = {
        "kernelspec": {
            "display_name": "Python 3",
            "language": "python",
            "name": "python3",
        },
        "language_info": {"name": "python", "version": "3.x"},
    }

    nbf.write(nb, NOTEBOOK / "Week5_MachineLearning.ipynb")

except Exception as exc:
    print("Notebook update warning:", exc)

# ---------------------------------------------------------------------
# PDF report generation.
# ---------------------------------------------------------------------
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm

styles = getSampleStyleSheet()
styles.add(
    ParagraphStyle(
        name="SmallSwift",
        parent=styles["BodyText"],
        fontSize=8,
        leading=10,
    )
)
styles.add(
    ParagraphStyle(
        name="TinySwift",
        parent=styles["BodyText"],
        fontSize=7,
        leading=8,
    )
)


def para(text, style="BodyText"):
    import xml.sax.saxutils as sax

    return Paragraph(sax.escape(str(text)), styles[style])


def table_for(frame, columns, widths=None):
    rows = [[para(column, "TinySwift") for column in columns]]

    for _, row in frame.iterrows():
        values = []

        for column in columns:
            value = row.get(column, "")

            if value is None or (np.isscalar(value) and pd.isna(value)):
                value = ""
            elif isinstance(value, (float, np.floating)):
                value = f"{value:.3f}"

            values.append(para(value, "TinySwift"))

        rows.append(values)

    table = Table(
        rows,
        repeatRows=1,
        colWidths=widths,
        hAlign="LEFT",
    )

    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17324D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C2CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F0F4F8")],
                ),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    return table


selected = (
    comparison[comparison["Selected"]].copy()
    if "Selected" in comparison.columns
    else comparison.iloc[0:0].copy()
)


def make_technical_report():
    path = REPORTS / "Technical_Report.pdf"

    story = [
        para("SwiftLog Technical Report", "Title"),
        para(
            "Machine Learning Track | Mission ML05: Intelligent Logistics "
            "Prediction & Supply Chain Optimization",
            "Heading2",
        ),
        Spacer(1, 8),
    ]

    sections = [
        (
            "1. Purpose",
            "This report documents the machine-learning workflow for four "
            "logistics tasks using the supplied SwiftLog educational datasets. "
            "The workflow includes data quality review, feature engineering, "
            "candidate algorithm comparison, held-out evaluation, limited "
            "hyperparameter tuning, permutation importance, saved model "
            "bundles, and a Streamlit prototype.",
        ),
        (
            "2. Data preparation",
            "Source tables were loaded from Data/Original_Datasets. "
            "Processed copies were reviewed for row counts, duplicates, and "
            "missing cells. Related order, item, product, customer, shipment, "
            "warehouse, and return tables were joined where required.",
        ),
        (
            "3. Targets and feature design",
            "Delivery delay: positive when DelayDays > 0; DelayDays and "
            "post-event delivery status are excluded from predictors. "
            "Product return: positive when an OrderID appears in Returns. "
            "Demand forecasting: weekly product quantity is the target, "
            "with lagged prior-week demand and rolling history as predictors. "
            "Customer value: next-month revenue is the target, with features "
            "summarizing the preceding three months of activity.",
        ),
        (
            "4. Evaluation design",
            "Classification tasks report accuracy, precision, recall, F1, "
            "and ROC-AUC where available. Demand and customer-value tasks "
            "report MAE, RMSE, and R-squared where available. The demand "
            "and customer-value tasks use chronological train/test splits. "
            "Results should be treated as experimental.",
        ),
        (
            "5. Model comparison",
            "The following table reports the available candidate algorithm "
            "metrics. Selected marks the model stored in the model bundle.",
        ),
    ]

    for heading, text in sections:
        story.extend(
            [
                para(heading, "Heading2"),
                para(text),
                Spacer(1, 6),
            ]
        )

    comparison_columns = [
        "Problem",
        "Algorithm",
        "Selected",
        "Tuned",
        "F1",
        "ROC_AUC",
        "MAE",
        "RMSE",
        "R2",
        "Accuracy",
    ]
    available_columns = [
        column for column in comparison_columns
        if column in comparison.columns
    ]
    story.append(table_for(comparison, available_columns))

    story.extend(
        [
            Spacer(1, 8),
            para("6. Interpretation", "Heading2"),
            para(
                "The delivery-delay target is imbalanced, and its ROC-AUC "
                "is close to 0.5. Its F1 score should therefore not be "
                "interpreted as strong ability to distinguish delayed from "
                "on-time shipments. The return classifier also shows weak "
                "discrimination. Negative held-out R-squared for demand and "
                "customer-value models indicates worse performance than a "
                "simple mean baseline on the test periods. These results "
                "are reported transparently rather than presented as "
                "successful forecasts."
            ),
            para("7. Feature importance and artefacts", "Heading2"),
            para(
                "Permutation importance is saved in "
                "Models/Permutation_Importance.csv and visualized in "
                "Models/Feature_Importance.png when available. Classification "
                "confusion matrices are saved under Models/. The selected "
                "pipelines are stored in Models/Best_Model.pkl. The "
                "Streamlit app loads this bundle for demonstrations."
            ),
            para("8. Assumptions and limitations", "Heading2"),
            para(
                "The delay target is defined from DelayDays because the "
                "supplied status field may disagree with the numeric delay "
                "field. A recorded return is treated as the return label. "
                "Supplier reliability is not modelled because the available "
                "tables do not directly link supplier IDs to shipments. "
                "Customer value is defined as next-month revenue rather than "
                "lifetime value. The data covers a limited historical period "
                "and does not establish production-grade predictive performance."
            ),
            para("9. Recommendations", "Heading2"),
            para(
                "Improve data capture for planned and actual delivery times; "
                "reconcile delivery status with delay days; capture supplier "
                "IDs on relevant records; maintain time-stamped inventory "
                "snapshots; define return labels consistently; collect more "
                "historical periods; evaluate models against simple baselines; "
                "and pilot predictions with human review before operational use."
            ),
            para("10. Reproduction", "Heading2"),
            para(
                "Install dependencies from App/requirements.txt, then run "
                "python run_swiftlog_project.py from the project root. The "
                "script regenerates processed datasets, model comparisons, "
                "saved model bundles, evaluation figures, reports, and the "
                "presentation. The notebook reviews the outputs."
            ),
        ]
    )

    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=1.3 * cm,
        leftMargin=1.3 * cm,
        topMargin=1.3 * cm,
        bottomMargin=1.3 * cm,
    ).build(story)


def make_exec_summary():
    path = REPORTS / "Executive_Summary.pdf"

    story = [
        para("SwiftLog Executive Summary", "Title"),
        para("Machine Learning Track | Week 5", "Heading2"),
        Spacer(1, 10),
        para("Objective", "Heading2"),
        para(
            "Build an educational machine-learning prototype to explore "
            "delivery delays, product returns, weekly product demand, and "
            "next-month customer revenue using the supplied SwiftLog datasets."
        ),
        para("What was completed", "Heading2"),
        para(
            "Reviewed source tables; engineered task-specific features; "
            "compared candidate algorithms; evaluated held-out predictions; "
            "saved selected model pipelines; generated comparison tables and "
            "evaluation visuals; and prepared a Streamlit demonstration app."
        ),
        para("Results at a glance", "Heading2"),
    ]

    summary = selected.copy()
    summary_columns = [
        "Problem",
        "Algorithm",
        "F1",
        "ROC_AUC",
        "MAE",
        "RMSE",
        "R2",
        "Accuracy",
    ]
    available_columns = [
        column for column in summary_columns
        if column in summary.columns
    ]

    story.append(table_for(summary, available_columns))

    story.extend(
        [
            Spacer(1, 8),
            para("What the results mean", "Heading2"),
            para(
                "The current models are not ready for operational use. "
                "The delay classifier has high F1 but ROC-AUC close to 0.5, "
                "suggesting weak class separation. The return classifier "
                "also shows weak discrimination. Demand and next-month "
                "customer-value regression have negative held-out R-squared, "
                "so they do not outperform a simple mean baseline on the "
                "test period."
            ),
            para("Recommended next steps", "Heading2"),
            para(
                "Collect richer historical data, verify target definitions, "
                "add operationally available predictors, compare against "
                "simple baselines, and validate on later time periods. "
                "Use the app only as a demonstration and keep human review "
                "in the decision process."
            ),
            Spacer(1, 8),
            para("Conclusion", "Heading2"),
            para(
                "The project delivers a documented prototype workflow and "
                "evaluation. The most important outcome is identifying the "
                "data and validation improvements needed before any business "
                "deployment."
            ),
        ]
    )

    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    ).build(story)


make_technical_report()
make_exec_summary()

# ---------------------------------------------------------------------
# Executive presentation.
# ---------------------------------------------------------------------
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_slide(title, body, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(245, 248, 252)

    header = slide.shapes.add_shape(
        1, 0, 0, prs.slide_width, Inches(1.15)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(23, 50, 77)
    header.line.fill.background()

    header_frame = header.text_frame
    header_frame.text = title
    header_frame.paragraphs[0].font.size = Pt(28)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    text_frame = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.55), Inches(11.8), Inches(4.9)
    ).text_frame
    text_frame.word_wrap = True
    text_frame.text = body
    text_frame.paragraphs[0].font.size = Pt(22)
    text_frame.paragraphs[0].font.color.rgb = RGBColor(35, 48, 62)

    if note:
        note_frame = slide.shapes.add_textbox(
            Inches(0.8), Inches(6.5), Inches(11.8), Inches(0.4)
        ).text_frame
        note_frame.text = note
        note_frame.paragraphs[0].font.size = Pt(11)
        note_frame.paragraphs[0].font.color.rgb = RGBColor(75, 88, 100)

    footer = slide.shapes.add_textbox(
        Inches(0.8), Inches(7.05), Inches(11.5), Inches(0.25)
    ).text_frame
    footer.text = f"SwiftLog ML05 | Week 5 | {len(prs.slides):02d}/15"
    footer.paragraphs[0].font.size = Pt(10)


slides = [
    (
        "SwiftLog Logistics",
        "Machine Learning for Logistics Prediction and Supply Chain Optimization",
        "Educational prototype; not a production deployment.",
    ),
    (
        "Business challenge",
        "Explore whether available logistics data can support earlier review "
        "of delivery delays, product returns, product demand, and customer revenue.",
    ),
    (
        "Project objectives",
        "Prepare the data, engineer useful features, compare algorithms, "
        "evaluate results, interpret limitations, and demonstrate predictions "
        "in a simple app.",
    ),
    (
        "Data used",
        "Customers, Orders, Order Items, Products, Shipments, Delivery "
        "Performance, Returns, Warehouses, Inventory, Employees, and Suppliers.",
    ),
    (
        "Data preparation",
        "Reviewed row counts, missing cells, and duplicates; standardized "
        "text; joined related tables; and saved cleaned and task-specific "
        "feature datasets.",
    ),
    (
        "Task 1: delivery delay",
        "Target: DelayDays greater than zero. Predictors include available "
        "warehouse, location, customer-city, and order-date information. "
        "DelayDays itself is excluded from the inputs.",
    ),
    (
        "Task 2: product returns",
        "Target: order appears in the Returns table. Predictors include "
        "available order timing, city, and item-level aggregates.",
    ),
    (
        "Task 3: weekly demand",
        "Target: product units per week. Features use prior-week demand, "
        "demand four weeks earlier, and a rolling average of prior weeks.",
    ),
    (
        "Task 4: customer value",
        "Target: customer revenue in the next month. Features summarize "
        "the preceding three months of purchasing activity.",
    ),
    (
        "Algorithms and evaluation",
        "Compared classification and regression algorithms using held-out "
        "evaluation. Classification uses accuracy, precision, recall, F1, "
        "and ROC-AUC where available; regression uses MAE, RMSE, and R-squared.",
    ),
    (
        "Delivery delay result",
        metric_line("Delivery Delay Classification"),
        "ROC-AUC near 0.5 means weak class discrimination despite a high F1 score.",
    ),
    (
        "Return prediction result",
        metric_line("Product Return Classification"),
        "ROC-AUC near 0.5 means the current features do not reliably separate "
        "returned and non-returned orders.",
    ),
    (
        "Regression results",
        metric_line("Product Demand Forecasting")
        + "\n"
        + metric_line("Customer Value Prediction (Next Month)"),
        "Negative R-squared means performance is worse than a simple mean "
        "baseline on the held-out period.",
    ),
    (
        "Business recommendations",
        "Improve delivery timestamps, reconcile status and delay fields, "
        "link suppliers to relevant records, collect inventory snapshots, "
        "and define consistent return labels.",
    ),
    (
        "Conclusion and next steps",
        "The workflow and demo app form an educational prototype. Collect "
        "more history, improve predictors, validate on future periods, and "
        "require human review before operational use.",
    ),
]

for slide_item in slides:
    add_slide(*slide_item)

prs.save(PRESENTATION / "Executive_Presentation.pptx")

# ---------------------------------------------------------------------
# README.
# ---------------------------------------------------------------------
readme = """# SwiftLog Logistics Analytics and AI Transformation

**Machine Learning Track | Mission ML05: Intelligent Logistics Prediction and Supply Chain Optimization**

## Project overview
This educational prototype explores four machine-learning tasks using the supplied SwiftLog datasets: delivery-delay classification, product-return classification, weekly product-demand forecasting, and next-month customer-revenue prediction. The results are experimental and are not validated for production use.

## Project structure
- `Data/Original_Datasets/`: supplied source CSV files.
- `Data/Processed_Datasets/`: cleaned tables, task features, quality summary, and feature data dictionary.
- `Data/Feature_Engineered_Dataset.csv`: order-level engineered dataset used for the return-prediction task.
- `Notebook/Week5_MachineLearning.ipynb`: data-quality review, model comparison, evaluation artefacts, and interpretation.
- `Models/`: saved pipelines, model comparison, permutation importance, and confusion matrices.
- `App/`: Streamlit demonstration app.
- `Reports/`: technical report and executive summary.
- `Presentation/`: executive presentation.

## Run the project
1. Install Python packages: `python -m pip install -r App/requirements.txt`.
2. From the project root, run `python run_swiftlog_project.py` to regenerate the processed data, models, evaluation artefacts, reports, notebook, and presentation.
3. Launch the app with `python -m streamlit run App/app.py`.

## Prediction tasks
1. **Delivery delay classification:** target is `DelayDays > 0`; `DelayDays` is not used as an input feature.
2. **Product return classification:** target is whether an order appears in the Returns table.
3. **Weekly product demand forecasting:** predicts product units per week using prior-week and prior-four-week demand features. The test period is later than the training period.
4. **Next-month customer value:** predicts customer revenue for a target month using purchasing activity from the preceding three months. The test period is chronological.

## Outputs
- `Models/Best_Model.pkl`
- `Models/Model_Comparison.xlsx`
- `Models/Permutation_Importance.csv`
- `Models/Feature_Importance.png`
- `Reports/Technical_Report.pdf`
- `Reports/Executive_Summary.pdf`
- `Presentation/Executive_Presentation.pptx`

## Reproducibility note
The included model bundle was generated with scikit-learn 1.8.0. Install the matching version from `App/requirements.txt` when possible; loading a model with a different scikit-learn version can produce compatibility warnings.

## Results and limitations
Model comparison metrics are recorded in `Models/Model_Comparison.xlsx`. The current classification ROC-AUC scores indicate weak class discrimination. The selected regression models have negative held-out R-squared, meaning they performed worse than a simple mean baseline on the test period. These results are reported transparently and should not be presented as reliable production forecasts.

The supplied supplier data does not directly link suppliers to shipments, so supplier reliability prediction is not claimed. Delivery status may disagree with numeric delay days; this project defines a delay as `DelayDays > 0`. Customer value means next-month revenue, not lifetime value. The data covers a limited period and lacks some operationally important predictors.

## Recommended next steps
Collect richer historical data, verify target definitions, improve predictors, compare against simple baselines, validate on later periods, monitor performance, and require human review before any operational use.
"""

(ROOT / "README.md").write_text(readme, encoding="utf-8")

print("Documentation, notebook, app warnings, README, and presentation refreshed.")
print("Updated outputs:")
print(f"  {NOTEBOOK / 'Week5_MachineLearning.ipynb'}")
print(f"  {REPORTS / 'Technical_Report.pdf'}")
print(f"  {REPORTS / 'Executive_Summary.pdf'}")
print(f"  {PRESENTATION / 'Executive_Presentation.pptx'}")